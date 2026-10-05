"""Working on a deck together with a co-author (optional pipeline).

    python coauthor.py inspect theirs.pptx [--slides 1 4 5]
        Style tokens of their deck: slide size, layouts, background colours, and for each text
        run the font, size, colour and box. Read the title, body, muted and citation lines off
        this dump to build a `Theme` for our inserted slides.

    python coauthor.py diff theirs.pptx ours.pptx
        Their new version vs our current build. Slides are matched by title text. For each of
        our slides found in their file: identical, text differs, notes differ. Then the slides
        that exist only in their file (their new or reworked material). Differences in our
        slides are usually our own later edits; anything else is an edit of theirs that we must
        keep or discuss before rebuilding.
"""
import argparse
import difflib
import sys
from collections import Counter

from pptx import Presentation
from pptx.util import Emu


def title_of(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t and len(t) < 90 and " — " not in t:
                return t
    return None


def texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def notes(slide):
    return slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""


def inspect(path, only=None):
    p = Presentation(path)
    print(f"size {Emu(p.slide_width).inches:.3f} x {Emu(p.slide_height).inches:.3f} in; "
          f"layouts {[l.name for l in p.slide_layouts]}; slides {len(p.slides)}")
    fonts = Counter()
    for i, s in enumerate(p.slides, 1):
        if only and i not in only:
            continue
        bg = s.background.fill
        try:
            bgc = str(bg.fore_color.rgb) if bg.type is not None else "-"
        except (AttributeError, TypeError):
            bgc = "?"
        print(f"=== slide {i} bg={bgc} title={title_of(s)!r}")
        for sh in s.shapes:
            box = f"({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f},{Emu(sh.width).inches:.2f},{Emu(sh.height).inches:.2f})"
            if sh.shape_type == 13:
                print(f"   picture {box}")
                continue
            if not sh.has_text_frame:
                continue
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if not r.text.strip():
                        continue
                    f = r.font
                    try:
                        col = str(f.color.rgb) if f.color and f.color.type is not None else "-"
                    except AttributeError:
                        col = "theme"
                    size = f.size.pt if f.size else None
                    fonts[(f.name, size, col)] += 1
                    print(f"   {box} {f.name} {size} {col} b={f.bold} | {r.text[:60]!r}")
    print("\nmost common (font, size, colour):")
    for k, n in fonts.most_common(12):
        print(f"   {n:4d}  {k}")


def diff(theirs, ours):
    pt, po = Presentation(theirs), Presentation(ours)
    our_by_title = {}
    for i, s in enumerate(po.slides, 1):
        our_by_title.setdefault(title_of(s), (i, s))
    matched = set()
    print("== our slides inside their file")
    for i, s in enumerate(pt.slides, 1):
        t = title_of(s)
        if t in our_by_title:
            oi, os_ = our_by_title[t]
            matched.add(i)
            same_t, same_n = texts(s) == texts(os_), notes(s) == notes(os_)
            print(f"  theirs#{i} ours#{oi} {t[:55]!r}: text {'same' if same_t else 'DIFF'}, notes {'same' if same_n else 'DIFF'}")
            if not same_t:
                for line in difflib.unified_diff(texts(s), texts(os_), lineterm="", n=0):
                    if line[:1] in "+-" and not line.startswith(("+++", "---")):
                        print("      ", line[:140].replace("\n", " / "))
    print("== only in their file")
    for i, s in enumerate(pt.slides, 1):
        if i not in matched:
            print(f"  theirs#{i} {title_of(s)!r}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("inspect")
    a1.add_argument("deck")
    a1.add_argument("--slides", type=int, nargs="*")
    a2 = sub.add_parser("diff")
    a2.add_argument("theirs")
    a2.add_argument("ours")
    a = ap.parse_args()
    inspect(a.deck, set(a.slides or [])) if a.cmd == "inspect" else diff(a.theirs, a.ours)
