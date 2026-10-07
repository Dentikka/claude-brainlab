"""Preflight checks on any .pptx, whoever built it: a date on the title slide, speaker notes on
content slides, duplicate titles, wording (no "not X but Y", no rating of the material, no
talk about the audience or delivery instructions on a slide), memes (about five per talk, one
per block at most). With --outline, first a table of contents: number, kind, layout, notes
length, title. With --camera-corner, also every shape inside that corner (a quarter of the
width by a quarter of the height by default) — only for a deck where the user decided to keep
the corner free for a camera inset.

    python deck_check.py deck.pptx
    python deck_check.py deck.pptx --outline
    python deck_check.py deck.pptx --camera-corner tr [--corner-frac 0.25]

Slides built by deckkit carry their kind in the slide name; for other decks the kind comes
from the layout name (title, divider, break and closing layouts need no notes).
Wording hits are hints for a human pass, not verdicts.
"""
import argparse
import os
import sys

from pptx import Presentation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import checks  # noqa: E402


def outline(prs) -> None:
    for i, sl in enumerate(prs.slides, 1):
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
        layout = (sl.slide_layout.name or "")[:24]
        print(f"{i:3d}  {checks.slide_kind(sl):8s} {layout:24s} notes {len(notes):5d}  {checks.slide_title(sl)[:70]}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--outline", action="store_true")
    ap.add_argument("--camera-corner", choices=checks.CORNERS,
                    help="corner kept free for a camera inset (only if the user chose one)")
    ap.add_argument("--corner-frac", type=float, default=0.25)
    a = ap.parse_args()
    prs = Presentation(a.deck)
    if a.outline:
        outline(prs)
        print()
    msgs = checks.report(prs)
    if a.camera_corner:
        msgs += [f"s{i:02d} in the camera corner ({a.camera_corner}): {what}"
                 for i, what in checks.corner_hits(prs, a.camera_corner, a.corner_frac)]
    for m in msgs:
        print(m)
    print(f"checks: {len(msgs)} issue(s)" if msgs else "checks: clean")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
