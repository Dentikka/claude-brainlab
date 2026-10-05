"""Checks a deck must pass before it is shown, on any .pptx (ours, a co-author's, a template's).

- the first slide carries a date (and, for a deck built by `LabDeck`, speaker names);
- every content slide has speaker notes (title, dividers, memes, breaks, closing need none);
- no two content slides share a title (notes are keyed by title; the audience gets lost);
- wording: no "not X but Y" aphorisms, no rating of the material itself, no talk about the
  audience in the third person and no delivery instructions on a slide (they go to the notes).

The kind of a slide comes from the tag `deckkit:<kind>` that the builders put into the slide
name, else from the layout name (course template codes, default layouts), else from the look
(a picture over a third of the slide with little text is a meme), else "content".
"""
from __future__ import annotations

import re
from typing import Iterable, List, Optional, Tuple

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

__all__ = ["slide_kind", "slide_title", "slide_text", "has_date", "missing_notes", "duplicate_titles",
           "wording", "report", "TAG"]

TAG = "deckkit:"
SERVICE = {"title", "divider", "meme", "break", "closing"}
SERVICE_LAYOUT = re.compile(r"·\s*Title\s*·|Section divider|\bBreak\b|\bClosing\b|^Title Slide$|^Section Header$")

_MONTHS_EN = "January|February|March|April|May|June|July|August|September|October|November|December"
_MONTHS_RU = "января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря"
DATE = re.compile(
    r"\b\d{1,2}[./]\d{1,2}[./]\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b"
    rf"|\b\d{{1,2}}\s+({_MONTHS_EN}|{_MONTHS_RU})\s+\d{{4}}\b|\b({_MONTHS_EN})\s+\d{{1,2}},\s+\d{{4}}\b",
    re.IGNORECASE)

# (rule, pattern) — hints for a human pass, not verdicts
WORDING = [
    ("not X but Y", re.compile(r"\bnot\s+(?:just\s+|only\s+|merely\s+)?[^.;:!?\n]{1,60}?[,;—–]\s*but\b", re.I)),
    ("not X but Y", re.compile(r"\b(?:is|are|was|were)\s+not\s+[^.;:!?\n]{1,60}?[,;—–]\s*"
                               r"(?:it|they|this|that)(?:'s|\s+is|\s+are|\s+was|\s+were)\b", re.I)),
    ("not X but Y", re.compile(r"\b(?:isn't|aren't|wasn't)\s+[^.;:!?\n]{1,60}?[,;—–]\s*"
                               r"(?:it|they|this|that)(?:'s|\s+is|\s+are)\b", re.I)),
    ("not X but Y", re.compile(r"\bit'?s\s+not\s+about\b", re.I)),
    ("не X, а Y", re.compile(r"\bэто\s+не\s+[^.;:!?\n]{1,60}?[,—–]\s*(?:это|а)\b", re.I)),
    ("не X, а Y", re.compile(r"\bне\s+[^.;:!?\n]{1,40}?,\s*а\s", re.I)),
    ("не X, а Y", re.compile(r"\b(?:дело|вопрос)\s+не\s+в\b", re.I)),
    ("rates itself", re.compile(r"\b(?:key|main|central|most\s+important|strongest|most\s+convincing|crucial)\s+"
                                r"(?:slide|result|number|point|takeaway|insight)\s+of\s+(?:the|this)\s+"
                                r"(?:talk|seminar|lecture|block|deck|course)\b", re.I)),
    ("rates itself", re.compile(r"\b(?:highlight|centerpiece|centrepiece)\s+of\s+(?:the|this)\s+"
                                r"(?:talk|seminar|lecture|block)\b", re.I)),
    ("rates itself", re.compile(r"\bгвоздь\b|\bглавн\w*\s+результат\w*\s+(?:семинара|доклада|лекции|блока)\b"
                                r"|\bсамое\s+убедительное\b|\bубеждает\s+лучше\b", re.I)),
    ("talks about the audience", re.compile(r"\b(?:the\s+)?(?:students|audience|listeners)\s+"
                                            r"(?:will|should|can|have|see|saw|understand|learn)\b", re.I)),
    ("delivery instruction", re.compile(r"\b(?:worth\s+saying|say\s+(?:it\s+)?(?:out\s+loud|aloud)|"
                                        r"ask\s+the\s+(?:audience|room|students))\b", re.I)),
    ("talks about the audience", re.compile(r"\bстуденты\s+(?:увидят|видят|поймут|должны)\b|\bсказать\s+вслух\b"
                                            r"|\bспросить\s+(?:зал|аудиторию)\b", re.I)),
]


def _looks_like_meme(slide) -> bool:
    """A picture over a third of the slide and little text: a meme built without a tag."""
    prs = slide.part.package.presentation_part.presentation
    area = prs.slide_width * prs.slide_height
    big = any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE and sh.width * sh.height >= area / 3 for sh in slide.shapes)
    return big and len(slide_text(slide)) < 150


def slide_kind(slide) -> str:
    name = slide._element.cSld.get("name") or ""
    if name.startswith(TAG):
        return name[len(TAG):]
    if SERVICE_LAYOUT.search(slide.slide_layout.name or ""):
        return "service"
    if _looks_like_meme(slide):
        return "meme"
    return "content"


def _texts(slide) -> Iterable[Tuple[float, str, object]]:
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            yield (max(sizes) if sizes else 0.0), sh.text_frame.text.strip(), sh


def slide_title(slide) -> str:
    """The title placeholder if there is one, else the text set in the largest type."""
    for sh in slide.placeholders:
        if sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) and sh.has_text_frame:
            if sh.text_frame.text.strip():
                return " ".join(sh.text_frame.text.split())
    best = max(_texts(slide), key=lambda t: t[0], default=None)
    return " ".join(best[1].split()) if best else ""


def slide_text(slide) -> str:
    return "\n".join(t for _, t, _ in _texts(slide))


def has_date(slide) -> bool:
    return bool(DATE.search(slide_text(slide)))


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    tf = slide.notes_slide.notes_text_frame
    return tf.text.strip() if tf is not None else ""


def missing_notes(prs, only_tagged: bool = False) -> List[Tuple[int, str]]:
    out = []
    for i, sl in enumerate(prs.slides, 1):
        kind = slide_kind(sl)
        if i == 1 or kind in SERVICE or kind == "service":
            continue
        if only_tagged and not (sl._element.cSld.get("name") or "").startswith(TAG):
            continue
        if not _notes(sl):
            out.append((i, slide_title(sl)))
    return out


def duplicate_titles(prs, only_tagged: bool = False) -> List[Tuple[str, List[int]]]:
    seen = {}
    for i, sl in enumerate(prs.slides, 1):
        if slide_kind(sl) != "content" or (only_tagged and not (sl._element.cSld.get("name") or "").startswith(TAG)):
            continue
        seen.setdefault(slide_title(sl), []).append(i)
    return [(t, ix) for t, ix in seen.items() if t and len(ix) > 1]


def wording(prs, only_tagged: bool = False) -> List[Tuple[int, str, str]]:
    """Slide text only: delivery instructions are fine in the notes, where they belong.
    One hit per slide and rule."""
    out = []
    for i, sl in enumerate(prs.slides, 1):
        if only_tagged and not (sl._element.cSld.get("name") or "").startswith(TAG):
            continue
        text = " ".join(slide_text(sl).split())
        seen = set()
        for rule, rx in WORDING:
            m = rx.search(text)
            if m and rule not in seen:
                seen.add(rule)
                out.append((i, rule, text[max(0, m.start() - 20):m.end() + 20]))
    return out


def report(prs, only_tagged: bool = False, title_meta: Optional[dict] = None) -> List[str]:
    """Human-readable problems; an empty list means the deck passed."""
    msgs = []
    if title_meta is not None:                       # the builder knows what it was given
        if not title_meta.get("date"):
            msgs.append("title slide: no date of the talk")
        if not title_meta.get("speakers"):
            msgs.append("title slide: no speaker names")
    elif len(prs.slides) and not has_date(prs.slides[0]):
        msgs.append("title slide: no date of the talk")
    for i, title in missing_notes(prs, only_tagged):
        msgs.append(f"s{i:02d} no speaker notes: {title[:60]!r}")
    for title, ix in duplicate_titles(prs, only_tagged):
        msgs.append(f"same title on slides {', '.join(map(str, ix))}: {title[:60]!r}")
    for i, rule, snippet in wording(prs, only_tagged):
        msgs.append(f"s{i:02d} wording ({rule}): ...{snippet}...")
    return msgs
