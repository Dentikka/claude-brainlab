"""Style-agnostic primitives for building a .pptx deck as code with python-pptx.

A deck is regenerated from a build script every time; nothing is edited by hand in the
.pptx. Every helper takes a `Deck`, whose `theme` carries the colours, fonts and the
geometry of the title and citation lines, so the same builders serve the lab style and a
co-author's style.

Navigation: a numbered divider starts a block; every content slide after it carries the
block's label above the title, and with `deck.blocks` set the dividers show a row of pills
("you are here"). Pills, agenda cards and labels become clickable jumps at save time.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence, Tuple

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from . import checks

__all__ = ["Theme", "Deck", "rgb", "arxiv_abs", "arxiv_pdf", "nav_pills"]

logger = logging.getLogger(__name__)

Box = Tuple[float, float, float, float]  # x, y, w, h in inches

HLINK_EXT = "{A12FA001-AC4F-418D-AE19-62706E023703}"
HLINK_NS = "http://schemas.microsoft.com/office/drawing/2018/hyperlinkcolor"


def rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6)


def arxiv_abs(aid: str) -> str:
    return f"https://arxiv.org/abs/{aid}"


def arxiv_pdf(aid: str) -> str:
    return f"https://arxiv.org/pdf/{aid}"


def nav_pills(slide, x: float, y: float, labels: Sequence[str], current: int, *, font: str,
              colors: Dict[str, Tuple[Optional[str], str]], line: str, size: float = 9, h: float = 0.24,
              gap: float = 0.08, w: Optional[float] = None, anchor: str = "left") -> list:
    """'You are here': one pill per block, blocks before `current` (0-based) 'done', the current
    one 'here', the rest 'ahead'. `colors[state] = (fill or None, text)`; `line` outlines the
    pills without a fill. `anchor` says what x is: the left edge, the right edge or the centre
    of the row. Works on any slide, including a course-template divider. Returns the pill
    shapes (for clickable jumps)."""
    widths = [w or max(0.42, 0.075 * len(lab) * size / 9 + 0.22) for lab in labels]
    total = sum(widths) + gap * (len(labels) - 1)
    cx = {"left": x, "right": x - total, "center": x - total / 2}[anchor]
    shapes = []
    for i, (lab, pw) in enumerate(zip(labels, widths)):
        state = "done" if i < current else ("here" if i == current else "ahead")
        fill, ink = colors[state]
        sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(y), Inches(pw), Inches(h))
        sh.adjustments[0] = 0.5
        sh.shadow.inherit = False
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
            sh.line.fill.background()
        else:
            sh.fill.background()
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(0.75)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = lab
        r.font.name, r.font.size, r.font.bold = font, Pt(size), state == "here"
        r.font.color.rgb = rgb(ink)
        shapes.append(sh)
        cx += pw + gap
    return shapes


@dataclass(frozen=True)
class Theme:
    """Everything a builder needs to know about a style. Colours are hex without '#'."""
    name: str
    width: float
    height: float
    font: str
    title_font: str
    mono_font: str
    colors: Dict[str, str]          # bg, bg_dark, title, ink, muted, card, card_alt, accent, accent2, gold, on_dark, muted_dark
    title_box: Box
    title_size: int
    lead_size: int = 17
    body_size: int = 14
    small_size: int = 12
    cite_box: Box = (0.89, 6.86, 11.9, 0.23)
    logo: Optional[Tuple[str, Box]] = None        # (png, box) on light slides
    logo_dark: Optional[Tuple[str, Box]] = None   # (png, box) on dark slides
    margin: float = 0.89

    def c(self, key: str) -> str:
        return self.colors[key]


@dataclass
class Deck:
    prs: Presentation
    theme: Theme
    layout_index: int = 0
    notes: Dict[str, str] = field(default_factory=dict)  # slide title -> speaker notes
    blocks: List[str] = field(default_factory=list)      # short block names, in order: agenda and pills
    auto_kicker: bool = True                             # block label above the title of content slides

    def __post_init__(self):
        self._block: Optional[Tuple[int, str]] = None    # current block (number, name)
        self._targets: Dict[tuple, object] = {}          # ("block", n) / ("agenda",) -> slide
        self._links: List[tuple] = []                    # (slide, shape, target key)
        self._title_meta: Optional[dict] = None          # set by a title builder that knows names and date

    # ------------------------------------------------------------------ slides
    @classmethod
    def blank(cls, theme: Theme, template: Optional[str] = None) -> "Deck":
        prs = Presentation(template) if template else Presentation()
        if not template:
            prs.slide_width, prs.slide_height = Inches(theme.width), Inches(theme.height)
        idx = 6 if not template and len(prs.slide_layouts) > 6 else 0  # 6 = 'Blank' in the default master
        return cls(prs, theme, idx)

    def new_slide(self, dark: bool = False, bg: Optional[str] = None, logo: bool = True, kind: str = "content"):
        """`kind` (content, title, divider, meme, closing) is stored in the slide name, where
        the checks read it: only content slides must have speaker notes."""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[self.layout_index])
        s._element.cSld.set("name", checks.TAG + kind)
        for ph in list(s.placeholders):  # builders draw everything themselves
            ph._element.getparent().remove(ph._element)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb(bg or self.theme.c("bg_dark" if dark else "bg"))
        if logo:
            self.add_logo(s, dark)
        return s

    def mark(self, s, kind: str) -> None:
        """Re-label a slide's kind (e.g. a closing built on a content slide)."""
        s._element.cSld.set("name", checks.TAG + kind)

    def add_logo(self, s, dark: bool = False) -> None:
        spec = self.theme.logo_dark if dark else self.theme.logo
        if spec:
            path, (x, y, w, h) = spec
            s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))

    # -------------------------------------------------------------- navigation
    def enter_block(self, number: int, name: str, divider_slide=None) -> None:
        """Start block `number`: later content slides carry its label; the divider is the
        target of the jumps to this block."""
        self._block = (number, name)
        if divider_slide is not None:
            self._targets[("block", number)] = divider_slide

    def kicker_text(self) -> Optional[str]:
        if not self._block:
            return None
        n, name = self._block
        return f"{n:02d} · {name.upper()}"

    def add_kicker(self, s, label: Optional[str] = None):
        """Block label above the title (None = the current block; '' = none). Click -> agenda."""
        label = self.kicker_text() if label is None and self.auto_kicker else label
        if not label:
            return None
        t = self.theme
        x, y, w, _ = t.title_box
        k = t.width / 10
        box = self.text(s, x, y - 0.26 * k, w, 0.22 * k, label, size=round(9 * k), bold=True,
                        color=t.c("accent"), spc=200)
        self.link(s, box, ("agenda",))
        return box

    def pill_colors(self, dark: bool = True):
        """(colors, line) for nav_pills from the theme: on a dark slide or a light one."""
        c = self.theme.c
        if dark:
            return ({"done": (c("muted_dark"), c("bg_dark")), "here": (c("accent2"), c("ink")),
                     "ahead": (None, c("muted_dark"))}, c("muted_dark"))
        return ({"done": (c("muted"), c("bg")), "here": (c("accent"), c("bg")), "ahead": (None, c("muted"))},
                c("muted"))

    def block_pills(self, s, x: float, y: float, anchor: str = "left", dark: bool = True, size: float = 9):
        """Pills for `self.blocks` with the current block lit; each jumps to its block's divider."""
        if not self.blocks or not self._block:
            return []
        colors, line = self.pill_colors(dark)
        k = self.theme.width / 10
        shapes = nav_pills(s, x, y, [f"{i:02d}" for i in range(1, len(self.blocks) + 1)], self._block[0] - 1,
                           font=self.theme.font, colors=colors, line=line, size=round(size * k),
                           h=0.24 * k, gap=0.08 * k, w=0.42 * k, anchor=anchor)
        for i, sh in enumerate(shapes, 1):
            self.link(s, sh, ("block", i))
        return shapes

    def link(self, s, shape, key: tuple) -> None:
        """Make `shape` jump to the slide registered under `key` (resolved at save time)."""
        self._links.append((s, shape, key))

    def _resolve_links(self) -> None:
        present = {sl.slide_id for sl in self.prs.slides}
        for s, shape, key in self._links:
            target = self._targets.get(key)
            if target is None or target.slide_id not in present or s.slide_id not in present:
                continue
            if target.slide_id != s.slide_id:
                shape.click_action.target_slide = target

    # -------------------------------------------------------------- primitives
    def text(self, s, x, y, w, h, runs, size=None, color=None, font=None, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spc=None, link=None):
        """`runs` is a string, or a list of paragraphs, each a string or a list of (text, bold)."""
        t = self.theme
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        for i, para in enumerate(runs if isinstance(runs, list) else [runs]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            for txt, b in (para if isinstance(para, list) else [(para, bold)]):
                r = p.add_run()
                r.text = txt
                f = r.font
                f.name, f.size, f.bold = font or t.font, Pt(size or t.body_size), b
                f.color.rgb = rgb(color or t.c("ink"))
                if spc:
                    f._rPr.set("spc", str(spc))
                if link:
                    r.hyperlink.address = link
        return box

    def bullets(self, s, x, y, w, h, items: Sequence, size=None, color=None, gap=8):
        box = self.text(s, x, y, w, h, list(items), size=size, color=color)
        for p in box.text_frame.paragraphs:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.24)))
            pPr.set("indent", str(-Inches(0.24)))
            clr = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(clr, qn("a:srgbClr")).set("val", self.theme.c("title"))
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
            p.space_after = Pt(gap)
        return box

    def card(self, s, x, y, w, h, fill: str, radius=0.06, line: Optional[str] = None):
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if line:
            sh.line.color.rgb = rgb(line)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def picture(self, s, path: str, x, y, w, h):
        """Fit an image into the box, centred, keeping its aspect ratio."""
        iw, ih = Image.open(path).size
        k = min(Inches(w) / iw, Inches(h) / ih)
        pw, ph = int(iw * k), int(ih * k)
        return s.shapes.add_picture(path, Inches(x) + (Inches(w) - pw) // 2, Inches(y) + (Inches(h) - ph) // 2, pw, ph)

    def outlined(self, box, width_pt: float = 1.5, color: str = "000000"):
        """Meme lettering: a dark outline under every run (a:ln must precede the fill)."""
        for p in box.text_frame.paragraphs:
            for r in p.runs:
                ln = etree.SubElement(r.font._rPr, qn("a:ln"))
                ln.set("w", str(int(Pt(width_pt))))
                etree.SubElement(etree.SubElement(ln, qn("a:solidFill")), qn("a:srgbClr")).set("val", color)
                r.font._rPr.insert(0, ln)
        return box

    def cite(self, s, label: str, url: Optional[str] = None):
        """Citation line: 'Author et al., Venue Year — Title'. The title is an underlined link
        in the text colour (hlinkClr=tx), the author part stays plain."""
        t = self.theme
        x, y, w, h = t.cite_box
        head, sep, title = label.partition(" — ")
        if not url:
            return self.text(s, x, y, w, h, label, size=t.small_size, color=t.c("muted"))
        runs = [[(head + sep, False), (title, False)]] if sep else [[(label, False)]]
        box = self.text(s, x, y, w, h, runs, size=t.small_size, color=t.c("muted"))
        run = box.text_frame.paragraphs[0].runs[-1]
        run.hyperlink.address = url
        run.font.underline = True
        click = run.font._rPr.find(qn("a:hlinkClick"))
        ext = etree.SubElement(etree.SubElement(click, qn("a:extLst")), qn("a:ext"))
        ext.set("uri", HLINK_EXT)
        etree.SubElement(ext, "{%s}hlinkClr" % HLINK_NS, nsmap={"ahyp": HLINK_NS}).set("val", "tx")
        return box

    def retext(self, shape, value: str) -> None:
        """Replace a shape's text, keeping the formatting of its first run (for edits of someone else's slide)."""
        p = shape.text_frame.paragraphs[0]
        p.runs[0].text = value
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)

    # ------------------------------------------------------------ data blocks
    def table(self, s, rows: List[Sequence[str]], widths: Sequence[float], x=None, y=1.6, row_h=0.42, size=13):
        t = self.theme
        x = t.margin if x is None else x
        tab = s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                                 Inches(row_h * len(rows))).table
        tab.first_row, tab.horz_banding = True, False
        for c, w in enumerate(widths):
            tab.columns[c].width = Inches(w)
        for r, row in enumerate(rows):
            tab.rows[r].height = Inches(row_h)
            for c, val in enumerate(row):
                cell = tab.cell(r, c)
                cell.text = str(val)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(t.c("title") if r == 0 else (t.c("bg") if r % 2 else t.c("card")))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.name, run.font.size, run.font.bold = t.font, Pt(size), r == 0
                        run.font.color.rgb = rgb(t.c("bg") if r == 0 else (t.c("title") if c == 0 else t.c("ink")))
        return tab

    def _labels(self, series, points_values, decimals):
        """Explicit label text: PowerPoint formats number labels with the system locale (a comma
        on a Russian Windows), so labels with decimals are written as text."""
        for ser, values in zip(series, points_values):
            for i, v in enumerate(values):
                tf = ser.points[i].data_label.text_frame
                tf.text = v if isinstance(v, str) else f"{v:.{decimals}f}"
                tf.paragraphs[0].runs[0].font.size = Pt(12)
                ser.points[i].data_label.position = XL_LABEL_POSITION.OUTSIDE_END

    def bar_chart(self, s, x, y, w, h, categories, series: Sequence[Tuple[str, Sequence[float]]],
                  colors: Optional[Sequence[str]] = None, decimals: Optional[int] = None):
        """Clustered columns, one colour per series, value axis hidden, labels on top."""
        t = self.theme
        data = CategoryChartData()
        data.categories = categories
        for name, values in series:
            data.add_series(name, values)
        chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        chart.font.name, chart.font.size = t.font, Pt(12)
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
        chart.category_axis.format.line.color.rgb = rgb("D9CFD4")
        for ser, col in zip(chart.series, colors or (t.c("title"), t.c("accent"), t.c("gold"))):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = rgb(col)
            ser.data_labels.show_value = True
            ser.data_labels.number_format, ser.data_labels.number_format_is_linked = "0", False
            ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        if decimals is not None:
            self._labels(chart.series, [v for _, v in series], decimals)
        chart.plots[0].gap_width = 70
        return chart

    def hbar_chart(self, s, x, y, w, h, categories, values, colors, labels=None, fmt="0%"):
        """Horizontal bars, one series, a colour per bar; categories read top to bottom."""
        t = self.theme
        data = CategoryChartData()
        data.categories = categories
        data.add_series("v", values)
        chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
        chart.font.name, chart.font.size = t.font, Pt(12)
        chart.font.color.rgb = rgb(t.c("ink"))
        chart.has_legend, chart.has_title = False, False
        chart.value_axis.visible = False
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = 0, max(values) * (1.25 if labels else 1.15)
        chart.category_axis.reverse_order = True
        chart.category_axis.format.line.color.rgb = rgb("D9CFD4")
        ser = chart.series[0]
        for i, col in enumerate(colors):
            ser.points[i].format.fill.solid()
            ser.points[i].format.fill.fore_color.rgb = rgb(col)
        ser.data_labels.show_value = True
        ser.data_labels.number_format, ser.data_labels.number_format_is_linked = fmt, False
        ser.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ser.data_labels.font.size, ser.data_labels.font.bold = Pt(12), True
        if labels:
            self._labels([ser], [labels], 0)
        chart.plots[0].gap_width = 45
        return chart

    # --------------------------------------------------------------- memes
    def meme(self, img: str, x: float, y: float, w: float,
             labels: Sequence[Tuple[float, float, float, float, str, int, str]]):
        """A meme on its own slide: picture at (x, y), width w; labels in image fractions
        (fx, fy, fw, fh, text, size, 'white'|'dark'). White labels get a dark outline."""
        s = self.new_slide(kind="meme")
        iw, ih = Image.open(img).size
        h = w * ih / iw
        s.shapes.add_picture(img, Inches(x), Inches(y), Inches(w), Inches(h))
        for fx, fy, fw, fh, txt, size, style in labels:
            box = self.text(s, x + fx * w, y + fy * h, fw * w, fh * h, txt.split("\n"), size=size,
                            color="FFFFFF" if style == "white" else self.theme.c("ink"), bold=True,
                            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            if style == "white":
                self.outlined(box)
        return s

    # ------------------------------------------------------- notes and order
    def note(self, s, text: str):
        """Speaker notes written next to the slide's code: they survive a renamed title."""
        s.notes_slide.notes_text_frame.text = text.strip()
        return s

    def apply_notes(self, skip_ids: Optional[set] = None) -> int:
        """Speaker notes keyed by the exact title text of a slide; slides in skip_ids
        (e.g. a co-author's) keep their own notes."""
        n = 0
        for sl in self.prs.slides:
            if skip_ids and sl.slide_id in skip_ids:
                continue
            hit = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text in self.notes]
            if hit:
                sl.notes_slide.notes_text_frame.text = self.notes[hit[0]]
                n += 1
        return n

    def reorder(self, order: Sequence) -> None:
        """Put slides in `order`; slides not listed are dropped together with their part."""
        assert len({sl.slide_id for sl in order}) == len(order), "a slide is listed twice"
        lst = self.prs.slides._sldIdLst
        by_id = {el.get("id"): el for el in lst}
        keep = [by_id[str(sl.slide_id)] for sl in order]
        for el in list(lst):
            lst.remove(el)
            if el not in keep:
                self.prs.part.drop_rel(el.get(qn("r:id")))
        for el in keep:
            lst.append(el)

    # ----------------------------------------------------------------- save
    def preflight(self) -> List[str]:
        """Problems to fix before the deck is shown (see deckkit.checks): date and names on the
        title slide, speaker notes on our content slides, duplicate titles, wording. A
        co-author's slides (not built here) are not checked."""
        return checks.report(self.prs, only_tagged=True, title_meta=self._title_meta)

    def save(self, path: str, check: bool = True) -> List[str]:
        """Resolve the clickable jumps, report preflight problems as warnings, write the file."""
        self._resolve_links()
        problems = self.preflight() if check else []
        for msg in problems:
            logger.warning("preflight: %s", msg)
        self.prs.save(path)
        return problems
