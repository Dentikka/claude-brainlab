"""LLM-course style: decks built on the course template's own layouts.

The course template is a .potx (the course repo's `templates/presentation_clean.potx`; pass its
path to `open_course` or set DECKKIT_COURSE_TEMPLATE) whose layouts carry coded names
("C-01 · ...", "F-04 · Section divider · dark") and numbered placeholders. Unlike the lab style, slides are filled through placeholders, not drawn.

Two template defects are handled here:
- python-pptx refuses a .potx: the main content type says "template"; `potx_to_pptx` rewrites it;
- the template has no notes master with a body placeholder, so speaker notes silently vanish;
  `ensure_notes_master` copies the default python-pptx notes master in.
"""
from __future__ import annotations

import io
import logging
import os
import zipfile
from typing import Dict, List, Optional, Union

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

from . import checks

__all__ = ["potx_to_pptx", "ensure_notes_master", "open_course", "dump_layouts", "CourseDeck"]

logger = logging.getLogger(__name__)

COURSE_TEMPLATE_ENV = "DECKKIT_COURSE_TEMPLATE"   # path to the course .potx when none is passed


def potx_to_pptx(src: str, dst: str) -> str:
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zo.writestr(it, data)
    return dst


def ensure_notes_master(src: str, dst: str) -> str:
    """Replace notesMaster1.xml with the default one from python-pptx when the template's
    master has no body placeholder (symptom: notes_text_frame is None)."""
    prs = Presentation(src)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    ok = s.notes_slide.notes_text_frame is not None
    if ok:
        import shutil
        shutil.copy(src, dst)
        return dst
    buf = io.BytesIO()
    blank = Presentation()
    blank.slides.add_slide(blank.slide_layouts[6]).notes_slide  # the default master is created lazily
    blank.save(buf)
    with zipfile.ZipFile(buf) as zd:
        default_master = zd.read("ppt/notesMasters/notesMaster1.xml")
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "ppt/notesMasters/notesMaster1.xml":
                data = default_master
            zo.writestr(it, data)
    return dst


def open_course(workdir: str, template: Optional[str] = None) -> Presentation:
    """Template -> usable Presentation with working notes (intermediate files in workdir).
    The template is `template` or, when it is not given, the path in $DECKKIT_COURSE_TEMPLATE."""
    template = template or os.environ.get(COURSE_TEMPLATE_ENV)
    if not template:
        raise ValueError(f"no course template: pass template= or set {COURSE_TEMPLATE_ENV}")
    tmp = potx_to_pptx(template, os.path.join(workdir, "_course.pptx"))
    fixed = ensure_notes_master(tmp, os.path.join(workdir, "_course_notes.pptx"))
    return Presentation(fixed)


def dump_layouts(prs: Presentation) -> str:
    """Placeholder map of every layout: read it before writing a builder for a layout."""
    out = []
    for lay in prs.slide_masters[0].slide_layouts:
        out.append(f"== {lay.name}")
        for ph in lay.placeholders:
            f = ph.placeholder_format
            out.append(f"  idx={f.idx:2d} {ph.name!r:40s} x={Emu(ph.left).inches:.2f} y={Emu(ph.top).inches:.2f} "
                       f"w={Emu(ph.width).inches:.2f} h={Emu(ph.height).inches:.2f}")
    return "\n".join(out)


class CourseDeck:
    """Fill course layouts by code: deck.slide("C-01", {0: title, 3: lead}, {7: "fig.png"}, notes).

    Navigation: set `deck.kicker = "BLOCK 6 · BIASES IN THE GRPO LOSS"` at each divider, and
    every C- slide gets it in its kicker placeholder (idx 1) unless the call fills idx 1 itself.
    Pills on an F-04 divider: `deckkit.core.nav_pills` with the template's colours."""

    def __init__(self, prs: Presentation, kicker: Optional[str] = None):
        self.prs = prs
        self.layouts = {l.name.split(" ")[0]: l for l in prs.slide_masters[0].slide_layouts}
        self.kicker = kicker

    @staticmethod
    def _fill(tf, value: Union[str, list]) -> None:
        items = value if isinstance(value, list) else [value]
        tf.text = ""
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item

    @staticmethod
    def _picture(slide, ph, path: str) -> None:
        x, y, w, h = ph.left, ph.top, ph.width, ph.height
        iw, ih = Image.open(path).size
        k = min(w / iw, h / ih)
        pw, phh = int(iw * k), int(ih * k)
        slide.shapes.add_picture(path, Emu(x + (w - pw) // 2), Emu(y + (h - phh) // 2), Emu(pw), Emu(phh))
        ph._element.getparent().remove(ph._element)

    def slide(self, code: str, texts: Optional[Dict[int, Union[str, list]]] = None,
              pics: Optional[Dict[int, str]] = None, notes: Optional[str] = None, kind: Optional[str] = None):
        """Placeholders not given text or a picture are removed, so no empty prompt text
        ('Click to add title') survives into the deck. `kind` ('meme', 'break', ...) marks a
        slide that needs no speaker notes when the layout name does not say so."""
        s = self.prs.slides.add_slide(self.layouts[code])
        texts, pics = dict(texts or {}), pics or {}
        if self.kicker and code.startswith("C-") and 1 not in texts and 1 not in pics:
            texts[1] = self.kicker
        if kind:
            s._element.cSld.set("name", checks.TAG + kind)
        for ph in list(s.placeholders):
            idx = ph.placeholder_format.idx
            if idx in pics:
                self._picture(s, ph, pics[idx])
            elif idx in texts:
                self._fill(ph.text_frame, texts[idx])
            else:
                ph._element.getparent().remove(ph._element)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    @staticmethod
    def note(slide, text: str):
        """Speaker notes next to the slide's code."""
        slide.notes_slide.notes_text_frame.text = text.strip()
        return slide

    def save(self, path: str, check: bool = True) -> List[str]:
        """Write the deck; report the preflight problems of deckkit.checks as warnings."""
        problems = checks.report(self.prs) if check else []
        for msg in problems:
            logger.warning("preflight: %s", msg)
        self.prs.save(path)
        return problems
