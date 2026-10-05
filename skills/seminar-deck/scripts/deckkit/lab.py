"""BRAIn lab house style (template "BRAIN LAB BASE"), the default for lab talks.

Tokens and geometry follow references/lab-brand.md (official template: 16:9, 10 x 5.625 in,
Trebuchet MS, maroon/cream/blue/gold). Signature elements: the card with a blue capsule tab,
the molecule "B" mark in the footer, oversized section numbers, molecule accent on the
title and closing slides, page numbers on content slides only.
"""
from __future__ import annotations

import os
from typing import Optional, Sequence, Tuple

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .builders import Builders
from .core import Theme, rgb

__all__ = ["LAB", "LabDeck", "ASSETS"]

ASSETS = os.path.join(os.path.dirname(__file__), "..", "..", "assets", "lab")
W, H = 10.0, 5.625
MARK_H = 0.34
MARK_W = MARK_H * 700 / 779


def asset(name: str) -> str:
    return os.path.normpath(os.path.join(ASSETS, name))


LAB = Theme(
    name="brainlab",
    width=W, height=H,
    font="Trebuchet MS", title_font="Trebuchet MS", mono_font="Consolas",
    colors=dict(bg="FAF8F1", bg_dark="854C65", title="854C65", title_dark="5B3445", ink="412430",
                muted="8A6677", card="FBEDD6", card_alt="F0C987", accent="0681B2", accent2="9DE1FC",
                gold="F0C987", on_dark="FAF8F1", muted_dark="E4D2DB", deep="412430"),
    title_box=(0.55, 0.42, 8.9, 0.7), title_size=24,
    lead_size=14, body_size=12, small_size=10,
    cite_box=(0.95, 5.25, 7.6, 0.22),
    logo=(asset("mark-color.png"), (0.42, H - 0.18 - MARK_H, MARK_W, MARK_H)),
    logo_dark=(asset("mark-light.png"), (0.42, H - 0.18 - MARK_H, MARK_W, MARK_H)),
    margin=0.55,
)


class LabDeck(Builders):
    """Lab-style deck: Builders with the house title, numbering and signature cards."""

    font_k = 0.8

    def __init__(self, prs, theme=LAB, layout_index=0, notes=None):
        super().__init__(prs, theme, layout_index, notes or {})
        self._page = 0

    @classmethod
    def create(cls) -> "LabDeck":
        base = Builders.blank(LAB)
        return cls(base.prs, LAB, base.layout_index)

    # ------------------------------------------------------------- footer
    def page_number(self, s, dark: bool = False) -> None:
        self._page += 1
        self.text(s, W - 1.1, H - 0.5, 0.7, 0.3, str(self._page), size=11,
                  color=self.theme.c("muted_dark" if dark else "muted"), align=PP_ALIGN.RIGHT)

    def content(self, title: str, source: Optional[str] = None, url: Optional[str] = None, page: bool = True):
        t = self.theme
        s = self.new_slide()
        self.text(s, *t.title_box, title, size=t.title_size, color=t.c("title"), bold=True)
        if source:
            self.cite(s, source, url)
        if page:
            self.page_number(s)
        return s

    # ------------------------------------------------------ house elements
    def capsule(self, s, x, y, text: str, w: float = 1.6, h: float = 0.3):
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        pill.adjustments[0] = 0.5
        pill.fill.solid()
        pill.fill.fore_color.rgb = rgb(self.theme.c("accent2"))
        pill.line.color.rgb = rgb(self.theme.c("accent"))
        pill.line.width = Pt(1.25)
        pill.shadow.inherit = False
        self.text(s, x + 0.16, y, w - 0.28, h, text, size=11, bold=True, color="0A3A52", anchor=MSO_ANCHOR.MIDDLE)

    def capsule_card(self, s, x, y, w, h, label: str = "", heading: str = "", body: str = "",
                     fill: str = "gold", sub: str = ""):
        """The signature card: rounded rectangle, blue capsule tab over the top-left corner."""
        t = self.theme
        col = {"gold": t.c("gold"), "cream": t.c("bg"), "warm": t.c("card")}.get(fill, fill)
        self.card(s, x, y, w, h, col, radius=0.08)
        if label:
            self.capsule(s, x + 0.22, y - 0.15, label, w=min(2.0, w - 0.5))
        top = y + (0.3 if label else 0.2)
        if heading:
            self.text(s, x + 0.26, top, w - 0.52, 0.36, heading, size=14, bold=True, color=t.c("deep"))
        if body:
            self.text(s, x + 0.26, top + (0.4 if heading else 0), w - 0.52, h - (top - y) - 0.45, body,
                      size=11.5, color=t.c("title_dark"))
        if sub:
            self.text(s, x + 0.26, y + h - 0.36, w - 0.52, 0.3, sub, size=9.5, color=t.c("title_dark"))

    def cards3(self, title, lead, cards: Sequence[Tuple[str, str, str, str]], foot="", source=None, url=None):
        t = self.theme
        s = self.content(title, source, url)
        self.text(s, 0.55, 1.2, 8.9, 0.35, lead, size=t.lead_size, color=t.c("title"))
        for i, (kick, head, body, sub) in enumerate(cards):
            self.capsule_card(s, 0.55 + 3.02 * i, 1.95, 2.84, 2.5, kick, head, body,
                              fill="gold" if i == 1 else "warm", sub=sub)
        if foot:
            self.text(s, 0.55, 4.65, 8.9, 0.3, foot, size=12, color=t.c("title"))
        return s

    def recap(self, title: str, items: Sequence[Tuple[str, str]], footer: str = ""):
        t = self.theme
        s = self.content(title)
        for i, (kick, body) in enumerate(items):
            self.capsule_card(s, 0.55, 1.35 + 1.12 * i, 8.9, 0.85, kick, body=body, fill="gold" if i == 0 else "warm")
        if footer:
            self.text(s, 0.55, 4.85, 8.9, 0.3, footer, size=12, color=t.c("title"))
        return s

    # ---------------------------------------------------------- dark slides
    def title_slide(self, title: str, subtitle: str = "", kicker: str = "", date: str = "", speakers: str = ""):
        t = self.theme
        s = self.new_slide(dark=True, logo=False)
        lw = 0.58 * 2500 / 779
        s.shapes.add_picture(asset("logo-light.png"), Inches(0.55), Inches(0.5), Inches(lw), Inches(0.58))
        self.picture(s, asset("molecule-node.png"), 6.95, 1.4, 2.7, 2.9)
        if kicker:
            self.text(s, 0.6, 2.36, 5.95, 0.4, kicker.upper(), size=13, bold=True, color=t.c("accent2"), spc=200)
        self.text(s, 0.6, 2.74, 5.95, 1.6, title, size=32, bold=True, color=t.c("on_dark"))
        if subtitle:
            self.text(s, 0.62, 4.3, 5.95, 0.5, subtitle, size=14, color=t.c("muted_dark"))
        if date or speakers:
            self.text(s, 0.62, 4.95, 5.95, 0.3, "  ·  ".join(x for x in (date, speakers) if x), size=11,
                      color=t.c("accent2"))
        return s

    def divider(self, kicker: str, title: str, sub: str = "", number: Optional[int] = None):
        """Section slide: oversized number upper-left, title lower; kicker and sub optional."""
        t = self.theme
        s = self.new_slide(dark=True)
        if number is not None:
            self.text(s, 0.55, 0.55, 3.2, 1.6, f"{number:02d}", size=96, bold=True, color=t.c("deep"))
        if kicker:
            self.text(s, 0.6, 2.35, 8.7, 0.3, kicker.upper(), size=12, bold=True, color=t.c("accent2"), spc=200)
        self.text(s, 0.6, 2.7, 8.7, 1.1, title, size=32, bold=True, color=t.c("on_dark"))
        if sub:
            self.text(s, 0.6, 3.85, 8.0, 0.8, sub, size=15, color=t.c("muted_dark"))
        return s

    def closing(self, title: str = "Thank you for your attention!", lines: Sequence[str] = (),
                picture: Optional[str] = None):
        """Maroon closing; a meme `picture` replaces the molecule cluster when given."""
        t = self.theme
        s = self.new_slide(dark=True, logo=False)
        lw = 0.58 * 2500 / 779
        s.shapes.add_picture(asset("logo-light.png"), Inches(0.55), Inches(0.5), Inches(lw), Inches(0.58))
        self.text(s, 0.6, 2.2 if picture else 2.7, 5.4 if picture else 5.95, 1.0, title, size=30, bold=True,
                  color=t.c("on_dark"))
        if lines:
            self.text(s, 0.62, 3.6, 5.4, 1.2, list(lines), size=13, color=t.c("muted_dark"))
        if picture:
            self.picture(s, picture, 6.1, 1.1, 3.5, 3.9)
        else:
            self.picture(s, asset("molecule.png"), 6.9, 1.7, 2.8, 2.8)
        return s
